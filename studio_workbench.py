"""Multi-domain documents and local, versioned content/project storage."""
from contextlib import contextmanager
import html
import io
import json
import re
import sqlite3
from datetime import datetime,timezone
from pathlib import Path
from uuid import uuid4
from flask import jsonify,request,send_file
from work_templates import DOMAINS,TEMPLATES,DOMAIN_MAP,TEMPLATE_MAP,work_request


def now():return datetime.now(timezone.utc).isoformat()

def text(value,limit,label):
    if not isinstance(value,str) or len(value)>limit:raise ValueError(f'{label} sai định dạng hoặc vượt {limit} ký tự.')
    return value

class WorkStore:
    def __init__(self,base):
        self.base=Path(base);self.path=self.base/'workbench.sqlite3'
        with self.connect() as db:
            db.executescript('''
            CREATE TABLE IF NOT EXISTS records(id TEXT PRIMARY KEY,title TEXT NOT NULL,kind TEXT NOT NULL,domain TEXT NOT NULL,
                version INTEGER NOT NULL,data TEXT NOT NULL,updated TEXT NOT NULL,archived INTEGER NOT NULL DEFAULT 0);
            CREATE TABLE IF NOT EXISTS versions(record_id TEXT NOT NULL,version INTEGER NOT NULL,data TEXT NOT NULL,
                updated TEXT NOT NULL,PRIMARY KEY(record_id,version));
            CREATE TABLE IF NOT EXISTS projects(id TEXT PRIMARY KEY,title TEXT NOT NULL,version INTEGER NOT NULL,
                data TEXT NOT NULL,updated TEXT NOT NULL);
            ''')
    @contextmanager
    def connect(self):
        db=sqlite3.connect(self.path,timeout=15);db.row_factory=sqlite3.Row
        try:
            with db:yield db
        finally:db.close()
    def normalize(self,data):
        if not isinstance(data,dict):raise ValueError('Nội dung phải là JSON object.')
        title=text(data.get('title','Nội dung mới'),200,'Tiêu đề').strip()
        if not title:raise ValueError('Nhập tiêu đề nội dung.')
        kind=data.get('kind','document');domain=data.get('domain','general')
        if kind not in {'document','image','video','voice'} or domain not in DOMAIN_MAP:raise ValueError('Loại nội dung/lĩnh vực không hợp lệ.')
        body=text(data.get('body',''),50000,'Nội dung')
        if kind=='document' and not body.strip():raise ValueError('Chưa có nội dung văn bản.')
        clean={'title':title,'kind':kind,'domain':domain,'body':body,'files':[],
               'template':text(data.get('template',''),80,'Mẫu'),
               'brief':text(data.get('brief',''),10000,'Yêu cầu'),
               'source':text(data.get('source',''),50000,'Tài liệu nguồn'),
               'provider':text(data.get('provider',''),40,'Nhà cung cấp'),
               'model':text(data.get('model',''),100,'Model'),'language':text(data.get('language','Tiếng Việt'),100,'Ngôn ngữ'),
               'audience':text(data.get('audience','Người đọc phổ thông'),300,'Người đọc'),'tone':text(data.get('tone','Rõ ràng, chuyên nghiệp'),100,'Giọng văn')}
        files=data.get('files',[])
        if not isinstance(files,list) or len(files)>100:raise ValueError('Danh sách tệp không hợp lệ.')
        for item in files:
            if not isinstance(item,dict):raise ValueError('Tệp không hợp lệ.')
            location=item.get('location','assets');filename=item.get('id')
            if location not in {'assets','voice'} or not isinstance(filename,str) or not re.fullmatch(r'[a-zA-Z0-9_.-]+',filename):raise ValueError('Đường dẫn tệp không hợp lệ.')
            ext=Path(filename).suffix.lower()
            allowed={'.png','.jpg','.jpeg','.webp','.bmp','.mp4','.mov','.avi','.mkv','.mp3','.wav','.m4a','.ogg'}
            if ext not in allowed:raise ValueError('Loại tệp không được hỗ trợ.')
            directory=self.base/('studio_assets' if location=='assets' else 'v7_voice_output')
            if not (directory/filename).is_file():raise ValueError('Tệp đã bị xóa hoặc chưa được tạo.')
            filekind='image' if ext in {'.png','.jpg','.jpeg','.webp','.bmp'} else 'audio' if ext in {'.mp3','.wav','.m4a','.ogg'} else 'video'
            clean['files'].append({'id':filename,'name':text(item.get('name',filename),200,'Tên tệp'),
                                   'kind':filekind,'location':location,'url':('/assets/' if location=='assets' else '/audio/')+filename})
        if kind!='document' and not files:raise ValueError('Nội dung này cần ít nhất một tệp.')
        return clean
    def save(self,data):
        clean=self.normalize(data);id=data.get('id');expected=data.get('version');stamp=now()
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            if id:
                row=db.execute('SELECT version FROM records WHERE id=?',(id,)).fetchone()
                if not row:raise ValueError('Nội dung không còn trong thư viện.')
                if expected!=row['version']:raise ValueError('Nội dung đã có phiên bản mới. Mở lại trước khi lưu.')
                version=row['version']+1
            else:id=uuid4().hex;version=1
            encoded=json.dumps(clean,ensure_ascii=False)
            db.execute('INSERT INTO records(id,title,kind,domain,version,data,updated) VALUES(?,?,?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET title=excluded.title,kind=excluded.kind,domain=excluded.domain,version=excluded.version,data=excluded.data,updated=excluded.updated,archived=0',
                       (id,clean['title'],clean['kind'],clean['domain'],version,encoded,stamp))
            db.execute('INSERT INTO versions VALUES(?,?,?,?)',(id,version,encoded,stamp))
        return {'id':id,'version':version,'updated':stamp,**clean}
    def list(self,query='',kind='',domain=''):
        with self.connect() as db:
            rows=db.execute('SELECT id,title,kind,domain,version,updated FROM records WHERE archived=0 AND (title LIKE ? OR data LIKE ?) AND (?="" OR kind=?) AND (?="" OR domain=?) ORDER BY updated DESC LIMIT 200',
                            ('%'+query+'%','%'+query+'%',kind,kind,domain,domain)).fetchall()
        return [dict(row) for row in rows]
    def get(self,id,version=None):
        with self.connect() as db:
            current=db.execute('SELECT * FROM records WHERE id=?',(id,)).fetchone()
            if not current:raise ValueError('Không tìm thấy nội dung.')
            row=db.execute('SELECT * FROM versions WHERE record_id=? AND version=?',(id,version or current['version'])).fetchone()
            if not row:raise ValueError('Không tìm thấy phiên bản.')
        return {'id':id,'version':row['version'],'latest_version':current['version'],'updated':row['updated'],**json.loads(row['data'])}
    def versions(self,id):
        with self.connect() as db:
            return [dict(row) for row in db.execute('SELECT version,updated FROM versions WHERE record_id=? ORDER BY version DESC LIMIT 50',(id,))]
    def archive(self,id,value=True):
        with self.connect() as db:db.execute('UPDATE records SET archived=? WHERE id=?',(int(value),id))
    def save_project(self,data):
        if not isinstance(data,dict):raise ValueError('Dữ liệu dự án không hợp lệ.')
        id=data.get('id','')
        if not isinstance(id,str) or not re.fullmatch(r'[A-Za-z0-9_-]{1,64}',id):raise ValueError('Mã dự án gồm chữ, số, gạch ngang hoặc gạch dưới.')
        title=text(data.get('title',id),200,'Tên dự án');project=data.get('project')
        if not isinstance(project,dict):raise ValueError('Snapshot dự án không hợp lệ.')
        allowed={'script','srt','scenes','media','music','parts','settings','ratio','resolution','music_volume','workspace'}
        clean={k:v for k,v in project.items() if k in allowed}
        encoded=json.dumps(clean,ensure_ascii=False)
        if len(encoded.encode())>2*1024*1024:raise ValueError('Snapshot vượt 2 MB. Giảm nội dung trước khi lưu.')
        stamp=now()
        with self.connect() as db:
            db.execute('BEGIN IMMEDIATE')
            row=db.execute('SELECT version FROM projects WHERE id=?',(id,)).fetchone()
            if row and data.get('version')!=row['version']:raise ValueError('Dự án đã tồn tại hoặc đã đổi. Tải lại trước khi lưu.')
            version=row['version']+1 if row else 1
            db.execute('INSERT INTO projects VALUES(?,?,?,?,?) ON CONFLICT(id) DO UPDATE SET title=excluded.title,version=excluded.version,data=excluded.data,updated=excluded.updated',
                       (id,title,version,encoded,stamp))
        return {'id':id,'version':version,'title':title,'updated':stamp}
    def projects(self):
        with self.connect() as db:return [dict(row) for row in db.execute('SELECT id,title,version,updated FROM projects ORDER BY updated DESC LIMIT 100')]
    def project(self,id):
        with self.connect() as db:row=db.execute('SELECT * FROM projects WHERE id=?',(id,)).fetchone()
        if not row:raise ValueError('Không tìm thấy dự án trên máy này.')
        return {'id':id,'version':row['version'],'title':row['title'],'project':json.loads(row['data'])}


def import_document(file):
    filename=file.filename or '';suffix=Path(filename).suffix.lower()
    binary=file.read(10*1024*1024+1)
    if len(binary)>10*1024*1024:raise ValueError('Tài liệu tối đa 10 MB.')
    try:
        if suffix in {'.txt','.md'}:
            result=binary.decode('utf-8-sig')
        elif suffix=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(io.BytesIO(binary))
            if reader.is_encrypted:raise ValueError('PDF có mật khẩu. Hãy cung cấp bản không khóa.')
            if len(reader.pages)>100:raise ValueError('PDF tối đa 100 trang.')
            result='\n\n'.join(page.extract_text() or '' for page in reader.pages)
        elif suffix=='.docx':
            import zipfile
            with zipfile.ZipFile(io.BytesIO(binary)) as archive:
                if sum(info.file_size for info in archive.infolist())>40*1024*1024:raise ValueError('DOCX có dữ liệu giải nén quá lớn.')
            from docx import Document
            doc=Document(io.BytesIO(binary))
            result='\n'.join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                result+='\n'+'\n'.join(' | '.join(cell.text for cell in row.cells) for row in table.rows)
        else:raise ValueError('Hỗ trợ TXT, MD, PDF và DOCX.')
    except UnicodeError:raise ValueError('Tệp văn bản cần mã hóa UTF-8.') from None
    except ValueError:raise
    except Exception:raise ValueError('Không đọc được tài liệu. Kiểm tra định dạng tệp.') from None
    if not result.strip():raise ValueError('Không trích xuất được văn bản. PDF ảnh scan cần OCR trước khi nhập.')
    if len(result)>50000:raise ValueError('Tài liệu vượt 50.000 ký tự. Chia nhỏ tài liệu trước khi nhập.')
    return Path(filename).name,result.strip()


def export_docx(title,body):
    from docx import Document
    from docx.shared import Pt
    doc=Document();doc.add_heading(title,0)
    doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(11)
    body=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',body)
    for line in body.splitlines():
        heading=re.match(r'^(#{1,3})\s+(.*)',line)
        if heading:doc.add_heading(heading[2],len(heading[1]))
        elif re.match(r'^[-*]\s+',line):doc.add_paragraph(line[2:],style='List Bullet')
        elif re.match(r'^\d+\.\s+',line):doc.add_paragraph(re.sub(r'^\d+\.\s+','',line),style='List Number')
        else:doc.add_paragraph(line)
    buffer=io.BytesIO();doc.save(buffer);buffer.seek(0);return buffer


def register_workbench(app,base,service):
    store=WorkStore(base);app.extensions['workbench']=store
    def body():
        data=request.get_json()
        if not isinstance(data,dict):raise ValueError('Dữ liệu phải là JSON object.')
        return data
    @app.get('/api/work/templates')
    def templates():return jsonify(ok=True,domains=[{k:v for k,v in d.items() if k!='instruction'} for d in DOMAINS],templates=TEMPLATES)
    @app.post('/api/work/generate')
    def generate():
        data=body();prompt,system,domain,template=work_request(data)
        output,model=service.text(data.get('provider'),prompt,system,data.get('model') or None)
        return jsonify(ok=True,text=output,model=model,domain=domain,template=template)
    @app.post('/api/work/import')
    def import_file():
        file=request.files.get('file')
        if not file:raise ValueError('Chọn tài liệu để nhập.')
        name,content=import_document(file)
        return jsonify(ok=True,name=name,text=content)
    @app.post('/api/work/export')
    def export():
        data=body();title=text(data.get('title','Tài liệu'),200,'Tiêu đề');content=text(data.get('body',''),50000,'Nội dung')
        if not content.strip():raise ValueError('Chưa có nội dung để xuất.')
        format=data.get('format','docx')
        if format=='docx':return send_file(export_docx(title,content),as_attachment=True,download_name='document.docx',mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
        if format in {'txt','md'}:return send_file(io.BytesIO(content.encode('utf-8')),as_attachment=True,download_name='document.'+format,mimetype='text/plain; charset=utf-8')
        if format=='html':
            output=f'<!doctype html><html lang="vi"><meta charset="utf-8"><title>{html.escape(title)}</title><style>body{{font:16px Arial;max-width:800px;margin:40px auto;padding:20px;color:#153e46}}pre{{white-space:pre-wrap;overflow-wrap:anywhere;font:inherit;line-height:1.7}}@media print{{body{{margin:0}}}}</style><h1>{html.escape(title)}</h1><pre>{html.escape(content)}</pre></html>'
            return send_file(io.BytesIO(output.encode()),as_attachment=True,download_name='document.html',mimetype='text/html')
        raise ValueError('Định dạng xuất không hợp lệ.')
    @app.get('/api/library/records')
    def list_records():return jsonify(ok=True,records=store.list(request.args.get('q','')[:200],request.args.get('kind',''),request.args.get('domain','')))
    @app.post('/api/library/records')
    def save_record():return jsonify(ok=True,record=store.save(body()))
    @app.get('/api/library/records/<id>')
    def get_record(id):
        version=request.args.get('version',type=int)
        return jsonify(ok=True,record=store.get(id,version),versions=store.versions(id))
    @app.post('/api/library/records/<id>/archive')
    def archive(id):
        data=body();store.archive(id,bool(data.get('archived',True)));return jsonify(ok=True)
    @app.get('/api/work/projects')
    def list_projects():return jsonify(ok=True,projects=store.projects())
    @app.post('/api/work/projects')
    def save_project():return jsonify(ok=True,**store.save_project(body()))
    @app.get('/api/work/projects/<id>')
    def load_project(id):return jsonify(ok=True,**store.project(id))
    return store
